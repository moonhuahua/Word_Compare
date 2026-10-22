from pathlib import Path
from io import BytesIO
import json
import hashlib
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from PIL import Image, ImageDraw

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / '示例'
OUT.mkdir(exist_ok=True)
FIXTURES = ROOT / 'tests' / 'fixtures'
FIXTURES.mkdir(exist_ok=True)

def picture(color):
    image = Image.new('RGB', (320, 120), 'white')
    draw = ImageDraw.Draw(image)
    draw.rectangle((15, 25, 305, 95), fill=color)
    draw.rectangle((60, 40, 260, 80), outline='white', width=4)
    buf = BytesIO()
    image.save(buf, format='PNG')
    buf.seek(0)
    return buf

def create_demo(revised):
    doc = Document()
    normal = doc.styles['Normal']
    normal.font.name = 'Microsoft YaHei'
    normal.font.size = Pt(11)
    doc.add_heading('纯化水系统技术方案', 0)
    doc.add_paragraph('项目演示文件，仅用于验证 Word_Compare 文档对比功能。')
    doc.add_heading('一、设计参数', 1)
    doc.add_paragraph('系统设计产水量为 %s m³/h，运行方式为全自动。' % ('12' if revised else '10'))
    doc.add_paragraph('压缩空气压力为 0.6～0.8 MPa。')
    if not revised:
        doc.add_paragraph('旧方案配置手动旁通阀。')
    doc.add_paragraph('电源要求：380 V / 50 Hz。')
    if revised:
        doc.add_paragraph('新增在线电导率仪，并预留远程通讯接口。')
    doc.add_heading('二、设备配置', 1)
    table = doc.add_table(rows=1, cols=4)
    table.style = 'Table Grid'
    for cell, text in zip(table.rows[0].cells, ['设备名称', '规格参数', '数量', '备注']):
        cell.text = text
    rows = [('原水泵', 'Q=%s m³/h' % ('15' if revised else '12'), '2', '一用一备'),
            ('反渗透主机', '%s m³/h' % ('12' if revised else '10'), '1', '自动运行'),
            ('纯水箱', '10 m³', '1', '材质 304')]
    if revised:
        rows.append(('电导率仪', '0～2000 μS/cm', '1', '新增'))
    for values in rows:
        for cell, value in zip(table.add_row().cells, values):
            cell.text = value
    doc.add_heading('三、交付与维护', 1)
    paragraph = doc.add_paragraph()
    run = paragraph.add_run('该段文字内容不变，仅调整字体格式。')
    run.bold = revised
    if revised:
        run.font.color.rgb = RGBColor.from_string('087F72')
        run.font.size = Pt(13)
    doc.add_paragraph('操作员培训、现场调试和运行维护资料均包含在交付范围内。')
    doc.add_paragraph('水质异常时，应检查仪表并按操作规程处理。')
    doc.add_paragraph('方案示意图片：')
    doc.add_picture(picture('#087F72' if revised else '#3366AA'), width=Inches(2.8))
    doc.sections[0].header.paragraphs[0].text = '中标环保 · 技术方案%s' % ('修订版' if revised else '原版')
    doc.sections[0].footer.paragraphs[0].text = 'Word_Compare 演示文档'
    doc.save(OUT / ('示例_修订版.docx' if revised else '示例_原版.docx'))

create_demo(False)
create_demo(True)

def simple(name, text, bold=False, split=False):
    doc = Document()
    p = doc.add_paragraph()
    if split:
        p.add_run(text[:2]).bold = bold
        p.add_run(text[2:]).bold = bold
    else:
        p.add_run(text).bold = bold
    doc.save(FIXTURES / name)

simple('测试_格式A.docx', '相同的中文内容')
simple('测试_格式B.docx', '相同的中文内容', bold=True)
simple('测试_分段运行.docx', '相同的中文内容', split=True)
simple('测试_空白A.docx', 'ABC 中文文字')
simple('测试_空白B.docx', 'abc  中文文字')
(FIXTURES/'测试_无效.docx').write_bytes(b'not-a-zip-file')
manifest={str(p.relative_to(ROOT)): hashlib.sha256(p.read_bytes()).hexdigest() for folder in [OUT,FIXTURES] for p in folder.glob('*') if p.is_file()}
(ROOT/'验证结果'/'源文件哈希.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(manifest, ensure_ascii=False, indent=2))
